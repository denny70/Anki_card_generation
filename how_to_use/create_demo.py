"""Create demo input files for Anki Card Generator."""
from openpyxl import Workbook

# Demo 1: Vietnamese (minimal - only front words)
wb = Workbook()
ws = wb.active
ws.title = 'Cards'
ws.append(['front_card', 'back_card_meaning', 'back_card_sentence', 'back_card_sentence_meaning', 'back_card_opposite', 'back_card_opposite_meaning'])
words_vi = ['xin chào', 'cảm ơn', 'đẹp', 'nhanh', 'ăn']
for w in words_vi:
    ws.append([w, '', '', '', '', ''])
ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 20
ws.column_dimensions['C'].width = 30
ws.column_dimensions['D'].width = 30
ws.column_dimensions['E'].width = 15
ws.column_dimensions['F'].width = 20
wb.save('demo_vietnamese.xlsx')
print('Created: demo_vietnamese.xlsx (5 words, minimal input)')

# Demo 2: Korean (minimal)
wb2 = Workbook()
ws2 = wb2.active
ws2.title = 'Cards'
ws2.append(['front_card', 'back_card_meaning', 'back_card_sentence', 'back_card_sentence_meaning', 'back_card_opposite', 'back_card_opposite_meaning'])
words_ko = ['안녕하세요', '감사합니다', '빠르다', '느리다', '먹다']
for w in words_ko:
    ws2.append([w, '', '', '', '', ''])
ws2.column_dimensions['A'].width = 15
ws2.column_dimensions['B'].width = 20
ws2.column_dimensions['C'].width = 30
ws2.column_dimensions['D'].width = 30
ws2.column_dimensions['E'].width = 15
ws2.column_dimensions['F'].width = 20
wb2.save('demo_korean.xlsx')
print('Created: demo_korean.xlsx (5 words, minimal input)')

# Demo 3: Vietnamese with some data pre-filled
wb3 = Workbook()
ws3 = wb3.active
ws3.title = 'Cards'
ws3.append(['front_card', 'back_card_meaning', 'back_card_sentence', 'back_card_sentence_meaning', 'back_card_opposite', 'back_card_opposite_meaning'])
ws3.append(['nóng', 'hot', '', '', 'lạnh', 'cold'])
ws3.append(['to', '', '', '', 'nhỏ', ''])
ws3.append(['sách', '', '', '', '', ''])
ws3.column_dimensions['A'].width = 15
ws3.column_dimensions['B'].width = 20
ws3.column_dimensions['C'].width = 30
ws3.column_dimensions['D'].width = 30
ws3.column_dimensions['E'].width = 15
ws3.column_dimensions['F'].width = 20
wb3.save('demo_partial_data.xlsx')
print('Created: demo_partial_data.xlsx (3 words, some data pre-filled)')

print('\nDone! Use these files with AnkiCardGenerator.')
print('Example: python gen_anki.py --input demo_vietnamese.xlsx --lang vi --deck "VN_demo"')
